from pathlib import Path
from docx import Document

src = Path('/home/ubuntu/projects/say-less-4c84aa09/Say_Less_Complete_Product_Documentation.docx')
out = Path('/home/ubuntu/review_sources/Say_Less_Complete_Product_Documentation.txt')
out.parent.mkdir(parents=True, exist_ok=True)
doc = Document(src)
lines = []
for para in doc.paragraphs:
    text = para.text.strip()
    if text:
        lines.append(text)
for table in doc.tables:
    lines.append('')
    for row in table.rows:
        cells = [cell.text.strip().replace('\n', ' / ') for cell in row.cells]
        lines.append(' | '.join(cells))
    lines.append('')
out.write_text('\n'.join(lines) + '\n', encoding='utf-8')
print(out)
print(f'paragraphs={len(doc.paragraphs)} tables={len(doc.tables)} lines={len(lines)}')
