from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "SAY_LESS_FULL_PRD.md"
OUTPUT = ROOT / "deliverables" / "SAY_LESS_FULL_PRD.docx"
LOGO = ROOT / "assets" / "images" / "icon.png"

NAVY = "111827"
LIME = "A6D900"
INK = "172033"
MUTED = "64748B"
LINE = "DDE2E8"
SOFT = "F4F6F8"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, **kwargs) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        if edge not in kwargs:
            continue
        edge_data = kwargs[edge]
        tag = "w:{}".format(edge)
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        for key in ["val", "sz", "space", "color"]:
            if key in edge_data:
                element.set(qn("w:{}".format(key)), str(edge_data[key]))


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    tr_pr.append(repeat)


def set_paragraph_spacing(paragraph, before=0, after=0, line=1.16) -> None:
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line


def add_field(paragraph, field: str) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = field
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(begin)
    run._r.append(instr)
    run._r.append(separate)
    run._r.append(end)


def style_document(doc: Document) -> None:
    section = doc.sections[0]
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.72)
    section.right_margin = Inches(0.72)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Aptos")
    normal.font.size = Pt(10.2)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.line_spacing = 1.18

    heading_sizes = {"Title": 30, "Heading 1": 18, "Heading 2": 14, "Heading 3": 11.5}
    for name, size in heading_sizes.items():
        style = styles[name]
        style.font.name = "Aptos Display" if name in ("Title", "Heading 1") else "Aptos"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), style.font.name)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(NAVY)
        style.paragraph_format.space_before = Pt(17 if name != "Title" else 0)
        style.paragraph_format.space_after = Pt(7)
        style.paragraph_format.keep_with_next = True

    header = section.header
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = hp.add_run("SAY LESS  /  PRODUCT REQUIREMENTS")
    run.font.name = "Aptos"
    run.font.size = Pt(8)
    run.font.bold = True
    run.font.color.rgb = RGBColor.from_string(MUTED)

    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = fp.add_run("CONFIDENTIAL PRODUCT BASELINE  ·  ")
    run.font.name = "Aptos"
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor.from_string(MUTED)
    add_field(fp, "PAGE")


def add_inline(paragraph, text: str, table=False) -> None:
    text = re.sub(r"\[([^\]]+)\]\([^\)]+\)", r"\1", text)
    tokens = re.split(r"(\*\*.*?\*\*|`.*?`|\*[^*]+\*)", text)
    for token in tokens:
        if not token:
            continue
        run = paragraph.add_run()
        if token.startswith("**") and token.endswith("**"):
            run.text = token[2:-2]
            run.bold = True
        elif token.startswith("`") and token.endswith("`"):
            run.text = token[1:-1]
            run.font.name = "Aptos Mono"
            run.font.size = Pt(8.5 if table else 9)
            run.font.color.rgb = RGBColor.from_string(MUTED)
        elif token.startswith("*") and token.endswith("*"):
            run.text = token[1:-1]
            run.italic = True
        else:
            run.text = token
        if table:
            run.font.size = Pt(8.3)


def parse_table_row(line: str) -> list[str]:
    return [part.strip() for part in line.strip().strip("|").split("|")]


def add_table(doc: Document, lines: list[str]) -> None:
    rows = [parse_table_row(line) for line in lines if "|" in line]
    if len(rows) < 2:
        return
    header = rows[0]
    data = rows[2:]
    columns = len(header)
    wide_table = columns >= 7
    if wide_table:
        landscape = doc.add_section(WD_SECTION.NEW_PAGE)
        landscape.orientation = WD_ORIENT.LANDSCAPE
        landscape.page_width, landscape.page_height = landscape.page_height, landscape.page_width
        landscape.top_margin = Inches(0.55)
        landscape.bottom_margin = Inches(0.55)
        landscape.left_margin = Inches(0.55)
        landscape.right_margin = Inches(0.55)
    table = doc.add_table(rows=1, cols=columns)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    table.style = "Table Grid"
    header_cells = table.rows[0].cells
    set_repeat_table_header(table.rows[0])
    for idx, text in enumerate(header):
        cell = header_cells[idx]
        set_cell_shading(cell, NAVY)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_border(cell, top={"val": "single", "sz": 4, "color": NAVY}, bottom={"val": "single", "sz": 4, "color": NAVY})
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        set_paragraph_spacing(p, before=2, after=2, line=1.0)
        add_inline(p, text, table=True)
        for run in p.runs:
            run.font.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)

    for row_index, row in enumerate(data):
        cells = table.add_row().cells
        for index in range(columns):
            cell = cells[index]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            if row_index % 2 == 1:
                set_cell_shading(cell, SOFT)
            set_cell_border(cell, bottom={"val": "single", "sz": 2, "color": LINE})
            p = cell.paragraphs[0]
            set_paragraph_spacing(p, before=2, after=2, line=1.04)
            add_inline(p, row[index] if index < len(row) else "", table=True)
    if wide_table:
        portrait = doc.add_section(WD_SECTION.NEW_PAGE)
        portrait.orientation = WD_ORIENT.PORTRAIT
        portrait.page_width, portrait.page_height = portrait.page_height, portrait.page_width
        portrait.top_margin = Inches(0.72)
        portrait.bottom_margin = Inches(0.65)
        portrait.left_margin = Inches(0.72)
        portrait.right_margin = Inches(0.72)
    else:
        doc.add_paragraph()


def add_cover(doc: Document, title: str, metadata: list[str], sections: list[str]) -> None:
    if LOGO.exists():
        logo_p = doc.add_paragraph()
        logo_p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        logo_p.add_run().add_picture(str(LOGO), width=Inches(0.55))
    tag = doc.add_paragraph()
    tag_run = tag.add_run("CONSENT-FIRST MEETING INTELLIGENCE")
    tag_run.font.name = "Aptos"
    tag_run.font.bold = True
    tag_run.font.size = Pt(9)
    tag_run.font.color.rgb = RGBColor.from_string("879800")
    set_paragraph_spacing(tag, before=24, after=8)

    title_p = doc.add_paragraph(style="Title")
    title_p.add_run(title.replace("# ", ""))
    set_paragraph_spacing(title_p, after=9, line=1.0)
    promise = doc.add_paragraph()
    promise_run = promise.add_run("Every meeting under the palm of your hands.")
    promise_run.font.size = Pt(15)
    promise_run.font.bold = True
    promise_run.font.color.rgb = RGBColor.from_string(LIME)
    set_paragraph_spacing(promise, after=25)

    meta_table = doc.add_table(rows=0, cols=2)
    meta_table.style = "Table Grid"
    for line in metadata:
        if not line.strip():
            continue
        clean = line.replace("**", "").strip()
        if ":" in clean:
            label, value = clean.split(":", 1)
        else:
            label, value = "Document", clean
        cells = meta_table.add_row().cells
        set_cell_shading(cells[0], NAVY)
        set_cell_shading(cells[1], SOFT)
        for cell in cells:
            set_cell_border(cell, bottom={"val": "single", "sz": 2, "color": LINE})
        p0 = cells[0].paragraphs[0]
        p1 = cells[1].paragraphs[0]
        set_paragraph_spacing(p0, before=4, after=4, line=1.0)
        set_paragraph_spacing(p1, before=4, after=4, line=1.0)
        p0.add_run(label.upper()).font.color.rgb = RGBColor(255, 255, 255)
        p0.runs[0].font.bold = True
        p0.runs[0].font.size = Pt(8)
        add_inline(p1, value)

    doc.add_paragraph()
    map_heading = doc.add_paragraph()
    run = map_heading.add_run("DOCUMENT MAP")
    run.font.bold = True
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor.from_string(MUTED)
    for section in sections:
        p = doc.add_paragraph(style="Normal")
        p.paragraph_format.left_indent = Inches(0.12)
        bullet = p.add_run("•  ")
        bullet.font.color.rgb = RGBColor.from_string(LIME)
        bullet.font.bold = True
        add_inline(p, section)
        set_paragraph_spacing(p, after=2)


def build_document() -> None:
    source_lines = SOURCE.read_text(encoding="utf-8").splitlines()
    title = next((line for line in source_lines if line.startswith("# ")), "Say Less — Full Product Requirements Document")
    title_index = source_lines.index(title)
    cover_end = next((i for i, line in enumerate(source_lines) if line.strip() == "---"), 9)
    metadata = source_lines[title_index + 1 : cover_end]
    document_map = [re.sub(r"^##\s+", "", line) for line in source_lines if line.startswith("## ")]

    doc = Document()
    style_document(doc)
    add_cover(doc, title, metadata, document_map)

    i = cover_end + 1
    while i < len(source_lines):
        line = source_lines[i]
        stripped = line.strip()
        if not stripped or stripped == "---":
            i += 1
            continue

        if stripped.startswith("|") and i + 1 < len(source_lines) and re.match(r"^\|\s*:?-{3,}", source_lines[i + 1].strip()):
            table_lines = [stripped, source_lines[i + 1].strip()]
            i += 2
            while i < len(source_lines) and source_lines[i].strip().startswith("|"):
                table_lines.append(source_lines[i].strip())
                i += 1
            add_table(doc, table_lines)
            continue

        heading = re.match(r"^(#{1,4})\s+(.+)$", stripped)
        if heading:
            level = len(heading.group(1))
            text = heading.group(2)
            if level == 2:
                doc.add_page_break()
            paragraph = doc.add_paragraph(style={1: "Title", 2: "Heading 1", 3: "Heading 2", 4: "Heading 3"}[level])
            add_inline(paragraph, text)
            i += 1
            continue

        numbered = re.match(r"^(\d+)\.\s+(.+)$", stripped)
        bullet = re.match(r"^[-*]\s+(.+)$", stripped)
        if numbered:
            p = doc.add_paragraph(style="Normal")
            p.paragraph_format.left_indent = Inches(0.28)
            p.paragraph_format.first_line_indent = Inches(-0.2)
            number = p.add_run(numbered.group(1) + ". ")
            number.font.bold = True
            number.font.color.rgb = RGBColor.from_string("879800")
            add_inline(p, numbered.group(2))
            set_paragraph_spacing(p, after=3)
            i += 1
            continue
        if bullet:
            p = doc.add_paragraph(style="Normal")
            p.paragraph_format.left_indent = Inches(0.28)
            p.paragraph_format.first_line_indent = Inches(-0.18)
            dot = p.add_run("• ")
            dot.font.bold = True
            dot.font.color.rgb = RGBColor.from_string(LIME)
            add_inline(p, bullet.group(1))
            set_paragraph_spacing(p, after=3)
            i += 1
            continue

        paragraph_lines = [stripped]
        i += 1
        while i < len(source_lines):
            candidate = source_lines[i].strip()
            if not candidate or candidate == "---" or candidate.startswith("#") or candidate.startswith("|"):
                break
            if re.match(r"^(\d+)\.\s+", candidate) or re.match(r"^[-*]\s+", candidate):
                break
            paragraph_lines.append(candidate)
            i += 1
        p = doc.add_paragraph(style="Normal")
        add_inline(p, " ".join(paragraph_lines))
        set_paragraph_spacing(p, after=7)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_document()
